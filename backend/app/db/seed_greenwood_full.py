"""Complete Ingestion and Seeder Script for Greenwood Public School (Dataset 1: Clean Baseline).

Loads all 55 Sections, 106 Teachers, 71 Rooms, 25 Subjects, and 587 Requirements
from Test/School_Timetable_Dataset_1_Clean.xlsx into ChronosEdu SQLite database.
Runs the OR-Tools CP-SAT scheduler to generate all 2,400 conflict-free lessons.
"""
from __future__ import annotations

import collections
import logging
import uuid
import openpyxl
from ortools.sat.python import cp_model
from sqlalchemy import delete, select

from app.core.security import hash_password
from app.db import models
from app.db.session import tenant_session
import chronos_engine
from chronos_engine import Rules, ClassSection, Subject, Teacher, Assignment, BlockedSlot, ProblemInput

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

from pathlib import Path
DEMO_SCHOOL_ID = uuid.UUID("a0000000-0000-0000-0000-000000000001")
BASE_DIR = Path(__file__).resolve().parent.parent.parent.parent
EXCEL_PATH = str(BASE_DIR / "Test" / "School_Timetable_Dataset_1_Clean.xlsx")


def ingest_greenwood_dataset_1():
    logger.info(f"Loading {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    sec_rows = {r[0]: r for r in list(wb["Sections"].iter_rows(values_only=True))[1:]}
    tea_rows = {r[0]: list(r) for r in list(wb["Teachers"].iter_rows(values_only=True))[1:]}
    req_rows = list(wb["Requirements"].iter_rows(values_only=True))[1:]
    room_rows = {r[0]: r for r in list(wb["Rooms"].iter_rows(values_only=True))[1:]}
    subj_rows = {r[0]: r for r in list(wb["Subjects"].iter_rows(values_only=True))[1:]}

    # Option A: PGT PHY & CHE max_w = 27
    for tid, t in tea_rows.items():
        if t[3] == "PGT" and t[4] in ("PHY", "CHE"):
            t[8] = 27

    logger.info("Computing optimal teacher assignments using CP-SAT...")
    m_assign = cp_model.CpModel()
    assign_vars = {}

    for i, r in enumerate(req_rows):
        sec_id, sub_id, periods, grp = r[1], r[4], r[6], r[9]
        ct_id = sec_rows[sec_id][8]

        if grp == "Class Teacher" or sub_id == "VE":
            eligible = [ct_id]
        elif grp == "NTT":
            eligible = [tid for tid, t in tea_rows.items() if t[3] == "NTT"]
        elif grp == "Specialist":
            eligible = [
                tid for tid, t in tea_rows.items()
                if t[3] == "Specialist" and (t[4] == sub_id or t[5] == sub_id)
            ]
        else:
            eligible = [
                tid for tid, t in tea_rows.items()
                if t[3] == grp and (t[4] == sub_id or t[5] == sub_id)
            ]

        vlist = []
        for tid in eligible:
            v = m_assign.NewBoolVar(f"a_{i}_{tid}")
            assign_vars[(i, tid)] = v
            vlist.append(v)
        m_assign.Add(sum(vlist) == 1)

    # Capacity constraints (VE is extra)
    for tid, t in tea_rows.items():
        max_w = t[8]
        t_expr = [
            r[6] * assign_vars[(i, tid)]
            for i, r in enumerate(req_rows)
            if (i, tid) in assign_vars and r[4] != "VE"
        ]
        if t_expr:
            m_assign.Add(sum(t_expr) <= max_w)

    # Objective: Prefer class teacher for that class section
    obj_prefs = []
    for i, r in enumerate(req_rows):
        sec_id, sub_id = r[1], r[4]
        ct_id = sec_rows[sec_id][8]
        if (i, ct_id) in assign_vars and r[9] != "Class Teacher":
            obj_prefs.append(10 * assign_vars[(i, ct_id)])

    if obj_prefs:
        m_assign.Maximize(sum(obj_prefs))

    s_assign = cp_model.CpSolver()
    st = s_assign.Solve(m_assign)
    if st not in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        raise RuntimeError("Teacher assignment optimization failed!")

    assignment_map = {}
    for i, r in enumerate(req_rows):
        sec_id, sub_id = r[1], r[4]
        for (idx, tid), var in assign_vars.items():
            if idx == i and s_assign.Value(var) == 1:
                assignment_map[(sec_id, sub_id)] = tid
                break

    logger.info(f"Assigned {len(assignment_map)} teaching requirements to 106 teachers.")

    # Generate Timetable Entries using chronos_engine
    working_days = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday")
    engine_rules = Rules(
        working_days=working_days,
        primary_periods=7,
        senior_periods=8,
        max_teacher_periods_per_day=6,
        activity_day="None",
        activity_periods=(),
        spread_subjects=False,
        smooth_teacher_workload=False,
    )

    engine_classes = []
    engine_blocked = []
    for sec_id, s in sec_rows.items():
        wing = s[4]
        tier = "primary" if wing in ("Pre-Primary", "Primary") else "senior"
        engine_classes.append(ClassSection(id=sec_id, name=sec_id, tier=tier, home_room=s[7] or ""))

        if wing == "Pre-Primary":
            for p in range(1, 8):
                engine_blocked.append(
                    BlockedSlot(scope="class", scope_id=sec_id, day="Saturday", period=p, label="No School")
                )
            for d in ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday"):
                engine_blocked.append(
                    BlockedSlot(scope="class", scope_id=sec_id, day=d, period=7, label="Early Dismissal")
                )

    engine_subjects = [
        Subject(id=sid, name=r[1], is_activity=(r[2] == "Activity"))
        for sid, r in subj_rows.items()
    ]

    teacher_subjects = collections.defaultdict(set)
    for (sec_id, sub_id), tid in assignment_map.items():
        teacher_subjects[tid].add(sub_id)

    engine_teachers = []
    for tid, t in tea_rows.items():
        is_ct = (t[13] is not None)
        actual_cap = t[8] + (2 if is_ct else 0)
        engine_teachers.append(
            Teacher(
                id=tid,
                name=t[1],
                subject_ids=frozenset(teacher_subjects[tid]),
                max_periods_per_day=t[9],
                max_periods_per_week=actual_cap,
                max_consecutive_periods=4,
            )
        )

    weekly_freq = {(r[1], r[4]): r[6] for r in req_rows}
    room_caps = {r[1]: (r[6] or 1) for rid, r in room_rows.items()}
    engine_assignments = tuple(
        Assignment(class_id=sec_id, subject_id=sub_id, teacher_id=tid)
        for (sec_id, sub_id), tid in assignment_map.items()
    )

    problem = ProblemInput(
        rules=engine_rules,
        classes=engine_classes,
        subjects=engine_subjects,
        teachers=engine_teachers,
        weekly_frequency=weekly_freq,
        assignments=engine_assignments,
        room_capacities=room_caps,
        blocked_slots=tuple(engine_blocked),
    )

    logger.info("Solving timetable for 2,400 periods...")
    solve_res = chronos_engine.solve(problem, time_limit_s=60.0)
    if not solve_res.feasible:
        raise RuntimeError("Solver returned infeasible on Dataset 1!")
    logger.info(f"Successfully generated {len(solve_res.entries)} timetable entries!")

    # Populate SQLite database
    logger.info("Persisting to ChronosEdu database...")
    with tenant_session(str(DEMO_SCHOOL_ID)) as session:
        # Wipe existing records for this school tenant
        session.execute(delete(models.TimetableEntry).where(models.TimetableEntry.school_id == DEMO_SCHOOL_ID))
        session.execute(delete(models.TeachingAssignment).where(models.TeachingAssignment.school_id == DEMO_SCHOOL_ID))
        session.execute(delete(models.SubjectRequirement).where(models.SubjectRequirement.school_id == DEMO_SCHOOL_ID))
        session.execute(delete(models.BlockedSlot).where(models.BlockedSlot.school_id == DEMO_SCHOOL_ID))
        session.execute(delete(models.PinnedLesson).where(models.PinnedLesson.school_id == DEMO_SCHOOL_ID))
        session.execute(delete(models.TeacherQualification).where(models.TeacherQualification.school_id == DEMO_SCHOOL_ID))
        session.execute(delete(models.Teacher).where(models.Teacher.school_id == DEMO_SCHOOL_ID))
        session.execute(delete(models.ClassSection).where(models.ClassSection.school_id == DEMO_SCHOOL_ID))
        session.execute(delete(models.Subject).where(models.Subject.school_id == DEMO_SCHOOL_ID))
        session.execute(delete(models.Room).where(models.Room.school_id == DEMO_SCHOOL_ID))
        session.execute(delete(models.GradeLevel).where(models.GradeLevel.school_id == DEMO_SCHOOL_ID))
        session.execute(delete(models.PeriodSlot).where(models.PeriodSlot.school_id == DEMO_SCHOOL_ID))
        session.execute(delete(models.TimetableVersion).where(models.TimetableVersion.school_id == DEMO_SCHOOL_ID))
        session.flush()

        # 1. School Tenant
        school = session.get(models.School, DEMO_SCHOOL_ID)
        if not school:
            school = models.School(id=DEMO_SCHOOL_ID, subdomain="greenwood")
            session.add(school)

        school.name = "Greenwood Public School"
        school.timezone = "Asia/Kolkata"
        school.term = "Academic Year 2026-27"
        school.working_days = list(working_days)
        school.published = True
        school.max_teacher_periods_per_day = 6
        school.activity_day = "None"
        school.activity_periods = []
        session.flush()

        # 2. Rooms
        room_id_map = {}
        for rid, r in room_rows.items():
            room_uuid = uuid.uuid4()
            room_id_map[rid] = room_uuid
            session.add(
                models.Room(
                    id=room_uuid,
                    school_id=DEMO_SCHOOL_ID,
                    name=r[1],
                    type=r[4].lower().replace(" ", "_"),
                    capacity=r[6] or 1,
                )
            )
        session.flush()

        # 3. GradeLevels
        # Group sections into GradeLevels (e.g. Nursery, LKG, UKG, Class 1..10, Class 11-PCM, etc.)
        grade_id_map = {} # class_name -> grade_uuid
        for sec_id, s in sec_rows.items():
            cls_name = s[1]
            wing = s[4]
            stream = s[5]
            grade_key = f"{cls_name} ({stream})" if stream != "General" else cls_name
            if grade_key not in grade_id_map:
                g_uuid = uuid.uuid4()
                grade_id_map[grade_key] = g_uuid
                tier = "primary" if wing in ("Pre-Primary", "Primary") else "senior"
                day_end = "12:20" if wing == "Pre-Primary" else ("13:30" if wing == "Primary" else "14:10")
                session.add(
                    models.GradeLevel(
                        id=g_uuid,
                        school_id=DEMO_SCHOOL_ID,
                        name=grade_key,
                        tier=tier,
                        day_end_time=day_end,
                    )
                )
        session.flush()

        # 4. ClassSections
        class_id_map = {} # sec_id -> class_uuid
        for sec_id, s in sec_rows.items():
            cls_name = s[1]
            stream = s[5]
            grade_key = f"{cls_name} ({stream})" if stream != "General" else cls_name
            g_uuid = grade_id_map[grade_key]
            home_room_uuid = room_id_map.get(s[7])
            c_uuid = uuid.uuid4()
            class_id_map[sec_id] = c_uuid
            session.add(
                models.ClassSection(
                    id=c_uuid,
                    school_id=DEMO_SCHOOL_ID,
                    grade_id=g_uuid,
                    section=sec_id,
                    home_room_id=home_room_uuid,
                )
            )
        session.flush()

        # 5. Subjects
        colors = [
            "bg-blue-50 text-blue-700 border-blue-200",
            "bg-indigo-50 text-indigo-700 border-indigo-200",
            "bg-emerald-50 text-emerald-700 border-emerald-200",
            "bg-amber-50 text-amber-700 border-amber-200",
            "bg-cyan-50 text-cyan-700 border-cyan-200",
            "bg-rose-50 text-rose-700 border-rose-200",
            "bg-purple-100 text-purple-800 border-purple-300",
            "bg-pink-100 text-pink-800 border-pink-300",
            "bg-lime-100 text-lime-800 border-lime-300",
            "bg-orange-50 text-orange-700 border-orange-200",
            "bg-teal-50 text-teal-700 border-teal-200",
        ]
        subj_id_map = {} # sub_code -> subj_uuid
        for idx, (sid, r) in enumerate(subj_rows.items()):
            sub_uuid = uuid.uuid4()
            subj_id_map[sid] = sub_uuid
            session.add(
                models.Subject(
                    id=sub_uuid,
                    school_id=DEMO_SCHOOL_ID,
                    name=r[1],
                    code=sid,
                    type="activity" if r[2] == "Activity" else "academic",
                    room_id=None,
                    requires_double_period=False,
                    color=colors[idx % len(colors)],
                )
            )
        session.flush()

        # 6. Teachers
        teacher_id_map = {} # tid -> teacher_uuid
        for tid, t in tea_rows.items():
            t_uuid = uuid.uuid4()
            teacher_id_map[tid] = t_uuid
            is_ct = (t[13] is not None)
            actual_cap = t[8] + (2 if is_ct else 0)
            email = f"{t[1].lower().replace(' ', '.')}@greenwood.edu"
            tea_model = models.Teacher(
                id=t_uuid,
                school_id=DEMO_SCHOOL_ID,
                name=t[1],
                email=email,
                phone="+91 98765 43210",
                max_periods_per_day=t[9],
                max_periods_per_week=actual_cap,
                max_consecutive_periods=4,
                unavailable_days=[],
            )
            session.add(tea_model)
        session.flush()

        # 7. Teacher Qualifications
        for tid, t in tea_rows.items():
            t_uuid = teacher_id_map[tid]
            subs = [t[4]]
            if t[5]:
                subs.append(t[5])
            for sid in subs:
                if sid in subj_id_map:
                    session.add(
                        models.TeacherQualification(
                            id=uuid.uuid4(),
                            school_id=DEMO_SCHOOL_ID,
                            teacher_id=t_uuid,
                            subject_id=subj_id_map[sid],
                            grade_id=None,
                        )
                    )
        session.flush()

        # 8. Subject Requirements
        # Add (grade_id, subject_id, weekly_frequency)
        seen_grade_req = set()
        for r in req_rows:
            sec_id, sub_id, periods = r[1], r[4], r[6]
            cls_name = sec_rows[sec_id][1]
            stream = sec_rows[sec_id][5]
            grade_key = f"{cls_name} ({stream})" if stream != "General" else cls_name
            g_uuid = grade_id_map[grade_key]
            sub_uuid = subj_id_map[sub_id]
            if (g_uuid, sub_uuid) not in seen_grade_req:
                seen_grade_req.add((g_uuid, sub_uuid))
                session.add(
                    models.SubjectRequirement(
                        id=uuid.uuid4(),
                        school_id=DEMO_SCHOOL_ID,
                        grade_id=g_uuid,
                        subject_id=sub_uuid,
                        weekly_frequency=periods,
                    )
                )
        session.flush()

        # 9. Teaching Assignments
        for (sec_id, sub_id), tid in assignment_map.items():
            session.add(
                models.TeachingAssignment(
                    id=uuid.uuid4(),
                    school_id=DEMO_SCHOOL_ID,
                    class_id=class_id_map[sec_id],
                    subject_id=subj_id_map[sub_id],
                    teacher_id=teacher_id_map[tid],
                )
            )
        session.flush()

        # 10. Period Slots (Bell Schedule)
        bell_times = [
            (1, "08:00", "08:40"),
            (2, "08:40", "09:20"),
            (3, "09:20", "10:00"),
            (4, "10:20", "11:00"),
            (5, "11:00", "11:40"),
            (6, "11:40", "12:20"),
            (7, "12:50", "13:30"),
            (8, "13:30", "14:10"),
        ]
        for p, st_t, end_t in bell_times:
            session.add(
                models.PeriodSlot(
                    id=uuid.uuid4(),
                    school_id=DEMO_SCHOOL_ID,
                    tier="senior",
                    kind="period",
                    period_number=p,
                    start_time=st_t,
                    end_time=end_t,
                    is_break=False,
                )
            )
            if p <= 7:
                session.add(
                    models.PeriodSlot(
                        id=uuid.uuid4(),
                        school_id=DEMO_SCHOOL_ID,
                        tier="primary",
                        kind="period",
                        period_number=p,
                        start_time=st_t,
                        end_time=end_t,
                        is_break=False,
                    )
                )
        session.flush()

        # 11. Timetable Version & Entries
        version = models.TimetableVersion(
            id=uuid.uuid4(),
            school_id=DEMO_SCHOOL_ID,
            version_number=1,
            state="published",
            notes="Greenwood Public School Master Timetable 2026-27 (2,400 periods)",
        )
        session.add(version)
        session.flush()

        # Entries
        for e in solve_res.entries:
            sec_id = e.class_id
            sub_id = e.subject_id
            tid = e.teacher_id
            home_room_code = sec_rows[sec_id][7]
            home_room_uuid = room_id_map.get(home_room_code)

            session.add(
                models.TimetableEntry(
                    id=uuid.uuid4(),
                    school_id=DEMO_SCHOOL_ID,
                    class_id=class_id_map[sec_id],
                    day=e.day,
                    period=e.period,
                    subject_id=subj_id_map[sub_id],
                    teacher_id=teacher_id_map[tid],
                    room_id=home_room_uuid,
                    is_manual_override=False,
                    version_id=version.id,
                )
            )
        session.flush()

        # 12. Admin User
        admin_email = "admin@greenwood.edu"
        existing_admin = session.scalar(
            select(models.AppUser).where(models.AppUser.school_id == DEMO_SCHOOL_ID, models.AppUser.email == admin_email)
        )
        if not existing_admin:
            session.add(
                models.AppUser(
                    id=uuid.uuid4(),
                    school_id=DEMO_SCHOOL_ID,
                    email=admin_email,
                    password_hash=hash_password("AdminPass123!"),
                    role="admin",
                )
            )
        session.flush()

    logger.info("Database ingestion and generation COMPLETE! All 2,400 entries successfully committed.")


if __name__ == "__main__":
    ingest_greenwood_dataset_1()
